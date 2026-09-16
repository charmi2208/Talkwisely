"""
Knowledge Base API Router.

POST   /api/v1/knowledge-base/upload   — Upload company PDF/DOCX/TXT knowledge document
GET    /api/v1/knowledge-base          — List uploaded knowledge documents
DELETE /api/v1/knowledge-base/{id}     — Delete document and purge vector embeddings
"""

import io
import os
import uuid
from pathlib import Path
from typing import Optional
import aiofiles
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.rag.vector_store import vector_store
from app.core.config import settings
from app.core.dependencies import get_current_user, get_db
from app.db.models import KnowledgeChunk, KnowledgeDocument

router = APIRouter(prefix="/knowledge-base", tags=["Knowledge Base"])

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100


def _extract_text(content: bytes, ext: str) -> str:
    """Pull plain text out of an uploaded document."""
    if ext in (".txt", ".md"):
        return content.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        import docx
        document = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in document.paragraphs)
    return ""


def _chunk_text(text: str) -> list[str]:
    step = CHUNK_SIZE - CHUNK_OVERLAP
    return [text[i:i + CHUNK_SIZE] for i in range(0, len(text), step) if text[i:i + CHUNK_SIZE].strip()]


async def _categories_for(db: AsyncSession, doc_ids: list[str]) -> dict[str, str]:
    """Category lives in the first chunk's metadata (documents have no category column)."""
    if not doc_ids:
        return {}
    res = await db.execute(
        select(KnowledgeChunk.document_id, KnowledgeChunk.extra_metadata).where(
            KnowledgeChunk.document_id.in_(doc_ids), KnowledgeChunk.chunk_index == 0
        )
    )
    return {doc_id: (meta or {}).get("category", "general") for doc_id, meta in res.all()}


class KnowledgeDocResponse(BaseModel):
    id: str
    title: str
    file_name: str
    file_type: str
    file_size_bytes: int
    category: Optional[str]
    chunk_count: int
    created_at: str

    class Config:
        from_attributes = True


@router.get("", response_model=list[KnowledgeDocResponse])
async def list_knowledge_documents(
    current_user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List knowledge base documents for current user's organization."""
    res = await db.execute(
        select(KnowledgeDocument)
        .where(KnowledgeDocument.organization_id == current_user.organization_id)
        .order_by(KnowledgeDocument.created_at.desc())
    )
    docs = res.scalars().all()
    categories = await _categories_for(db, [d.id for d in docs])

    return [
        KnowledgeDocResponse(
            id=d.id,
            title=d.title,
            file_name=d.file_name,
            file_type=d.file_type,
            file_size_bytes=d.file_size_bytes or 0,
            category=categories.get(d.id, "general"),
            chunk_count=d.chunk_count,
            created_at=d.created_at.isoformat(),
        )
        for d in docs
    ]


@router.post("/upload", response_model=KnowledgeDocResponse, status_code=status.HTTP_201_CREATED)
async def upload_knowledge_document(
    file: UploadFile = File(...),
    title: str = Form(...),
    category: Optional[str] = Form("general"),
    current_user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Upload a company knowledge document (PDF/DOCX/TXT) and index into Vector DB."""
    content_bytes = await file.read()
    file_size = len(content_bytes)

    if file_size == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    ext = Path(file.filename or "doc").suffix.lower()
    if ext not in (".pdf", ".docx", ".txt", ".md"):
        raise HTTPException(status_code=400, detail="Only PDF, DOCX, TXT, and MD files are supported")

    try:
        text_content = _extract_text(content_bytes, ext)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read this file. Check that it's a valid {ext.lstrip('.').upper()} document.",
        )
    chunks = _chunk_text(text_content)
    if not chunks:
        raise HTTPException(
            status_code=400,
            detail="No readable text found in this file. Scanned PDFs need OCR before upload.",
        )

    doc_id = str(uuid.uuid4())
    category = category or "general"
    save_dir = os.path.join(settings.storage_local_path, "kb", current_user.organization_id)
    os.makedirs(save_dir, exist_ok=True)
    file_path = os.path.join(save_dir, f"{doc_id}{ext}")

    async with aiofiles.open(file_path, "wb") as f:
        await f.write(content_bytes)

    db_doc = KnowledgeDocument(
        id=doc_id,
        organization_id=current_user.organization_id,
        uploaded_by=current_user.id,
        title=title.strip(),
        file_name=file.filename or "file",
        file_path=file_path,
        file_type=ext.replace(".", ""),
        file_size_bytes=file_size,
        status="indexed",
        chunk_count=len(chunks),
    )
    db.add(db_doc)

    chunk_ids = [f"{doc_id}_chunk_{idx}" for idx in range(len(chunks))]
    metadata = {
        "organization_id": current_user.organization_id,
        "document_id": doc_id,
        "title": db_doc.title,
        "category": category,
        "doc_type": "knowledge_base",
    }
    for idx, (chunk_id, chunk) in enumerate(zip(chunk_ids, chunks)):
        db.add(KnowledgeChunk(
            document_id=doc_id,
            organization_id=current_user.organization_id,
            chunk_index=idx,
            text=chunk,
            vector_id=chunk_id,
            extra_metadata={"category": category, "title": db_doc.title},
        ))

    try:
        await db.flush()
        await vector_store.add_documents(
            ids=chunk_ids,
            documents=chunks,
            metadatas=[{**metadata, "chunk_index": i} for i in range(len(chunks))],
        )
        await db.commit()
    except Exception:
        await db.rollback()
        await vector_store.delete_by_document(doc_id)
        if os.path.exists(file_path):
            os.remove(file_path)
        raise
    await db.refresh(db_doc)

    return KnowledgeDocResponse(
        id=db_doc.id,
        title=db_doc.title,
        file_name=db_doc.file_name,
        file_type=db_doc.file_type,
        file_size_bytes=db_doc.file_size_bytes or 0,
        category=category,
        chunk_count=db_doc.chunk_count,
        created_at=db_doc.created_at.isoformat(),
    )


@router.delete("/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_knowledge_document(
    document_id: str,
    current_user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete knowledge document and purge its vector embeddings."""
    res = await db.execute(
        select(KnowledgeDocument).where(
            KnowledgeDocument.id == document_id,
            KnowledgeDocument.organization_id == current_user.organization_id,
        )
    )
    doc = res.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    if doc.file_path and os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except Exception:
            pass

    await vector_store.delete_by_document(document_id)
    await db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id))
    await db.delete(doc)
    await db.commit()
